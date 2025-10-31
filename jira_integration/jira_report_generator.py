#!/usr/bin/env python3
"""
Jira Report Generator Script

Generates work reports from Jira tasks for a specific month
"""

import os
import sys
import calendar
import logging
from datetime import datetime, timedelta
from pathlib import Path
from atlassian import Jira
from docx import Document
from dotenv import load_dotenv

# Configure logging
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

# Load environment variables from project root
_project_root = Path(__file__).parent.parent.resolve()
load_dotenv(_project_root / 'config.env')

class JiraReportGenerator:
    def __init__(self):
        # Get project root directory (parent of jira_integration folder)
        self.project_root = Path(__file__).parent.parent.resolve()
        
        self.jira_server = os.getenv('JIRA_SERVER')
        self.jira_username = os.getenv('JIRA_USERNAME')
        self.jira_api_token = os.getenv('JIRA_API_TOKEN')
        self.report_author = os.getenv('REPORT_AUTHOR', 'Unknown User')
        self.jira_client = None
        
        if not all([self.jira_server, self.jira_username, self.jira_api_token]):
            raise ValueError("Missing Jira configuration. Check config.env file.")
        
        self._connect_to_jira()
    
    def _connect_to_jira(self):
        """Connect to Jira using API token authentication"""
        try:
            self.jira_client = Jira(
                url=self.jira_server,
                username=self.jira_username,
                password=self.jira_api_token,
                cloud=True
            )
            logger.info(f"Successfully connected to Jira: {self.jira_server}")
            logger.info(f"Jira client API root: {self.jira_client.resource_url('')}")
        except Exception as e:
            raise ConnectionError(f"Failed to connect to Jira: {e}")
    
    def parse_month_year(self, month_text, current_year=None):
        """Parse month name and return month number and year"""
        if current_year is None:
            current_year = datetime.now().year
            
        month_names = {
            'january': 1, 'jan': 1, 'janvar': 1,
            'february': 2, 'feb': 2, 'fevral': 2,
            'march': 3, 'mar': 3, 'mart': 3,
            'april': 4, 'apr': 4, 'aprel': 4,
            'may': 5, 'maj': 5,
            'june': 6, 'jun': 6, 'ijun': 6,
            'july': 7, 'jul': 7, 'ijul': 7,
            'august': 8, 'aug': 8, 'avgust': 8,
            'september': 9, 'sep': 9, 'sentjabr': 9,
            'october': 10, 'oct': 10, 'oktjabr': 10,
            'november': 11, 'nov': 11, 'nojabr': 11,
            'december': 12, 'dec': 12, 'dekabr': 12
        }
        
        month_lower = month_text.lower().strip()
        month_num = month_names.get(month_lower)
        
        if month_num is None:
            raise ValueError(f"Unknown month: {month_text}")
        
        return month_num, current_year
    
    def get_month_date_range(self, month_num, year):
        """Get start and end dates for the given month"""
        first_day = datetime(year, month_num, 1)
        last_day_num = calendar.monthrange(year, month_num)[1]
        last_day = datetime(year, month_num, last_day_num, 23, 59, 59)
        
        return first_day, last_day
    
    def _jql_search_v3(self, jql, limit=200):
        """Execute JQL search using API v3 enhanced search endpoint"""
        logger.info(f"Executing JQL search with API v3: {jql}")
        
        # Build the full API v3 URL manually - using the NEW /search/jql endpoint
        base_url = self.jira_server.rstrip('/')
        full_url = f"{base_url}/rest/api/3/search/jql"
        
        params = {
            'jql': jql,
            'maxResults': limit,
            'fields': ['summary', 'key']
        }
        
        try:
            logger.info(f"API v3 URL: {full_url}")
            logger.info(f"Query params: {params}")
            
            # Use the Jira client's session directly to bypass the resource_url() method
            response = self.jira_client.session.get(
                full_url,
                params=params,
                auth=(self.jira_username, self.jira_api_token)
            )
            
            logger.info(f"Response status code: {response.status_code}")
            
            if response.status_code != 200:
                error_msg = response.text
                logger.error(f"API returned error: {error_msg}")
                raise Exception(f"API request failed with status {response.status_code}: {error_msg}")
            
            result = response.json()
            logger.info(f"API v3 response: success, issues count: {len(result.get('issues', []))}")
            
            return result
            
        except Exception as e:
            logger.error(f"API v3 request failed: {e}")
            raise
    
    def get_user_tasks_for_month(self, month_num, year):
        """Get all tasks where current user participated in the given month"""
        first_day, last_day = self.get_month_date_range(month_num, year)
        
        # Format dates for JQL
        start_date = first_day.strftime('%Y-%m-%d')
        end_date = last_day.strftime('%Y-%m-%d')
        
        # JQL to find issues where user was assignee or reporter
        # Use specific username instead of currentUser() for better compatibility
        username = self.jira_username
        project_key = os.getenv("JIRA_PROJECT_KEY", "RGS")
        
        # Statuses to exclude (configurable via environment or use default)
        # Common excluded statuses: New, Postponed, Backlog, etc.
        excluded_statuses = os.getenv("JIRA_EXCLUDED_STATUSES", "New,Postponed,Backlog")
        
        if excluded_statuses and excluded_statuses.strip():
            # Build status exclusion clause
            statuses_list = [f'"{s.strip()}"' for s in excluded_statuses.split(',')]
            status_exclude = f'AND status NOT IN ({", ".join(statuses_list)})'
        else:
            # No status exclusion
            status_exclude = ''
        
        jql_queries = [
            # Tasks assigned to user updated in the period
            f'assignee = "{username}" AND updated >= "{start_date}" AND updated <= "{end_date}" {status_exclude}',
            # Tasks created by user
            f'reporter = "{username}" AND created >= "{start_date}" AND created <= "{end_date}" {status_exclude}'
        ]
        
        print(f"Using JQL queries with username: {username}")
        print(f"Date range: {start_date} to {end_date}")
        print(f"Project key: {os.getenv('JIRA_PROJECT_KEY', 'RGS')}")
        
        all_issues_dict = {}  # Use dict to avoid duplicates by key
        
        for i, jql in enumerate(jql_queries, 1):
            try:
                try:
                    logger.info(f"\nExecuting JQL query {i}/{len(jql_queries)}: {jql}")
                    print(f"\nExecuting JQL query {i}/{len(jql_queries)}: {jql}")
                except UnicodeEncodeError:
                    logger.info(f"\nExecuting JQL query {i}/{len(jql_queries)}")
                    print(f"\nExecuting JQL query {i}/{len(jql_queries)}")
                
                # Use API v3 search method
                issues_data = self._jql_search_v3(jql, limit=200)
                issues = issues_data.get('issues', [])
                
                logger.info(f"Raw response: {len(issues)} issues")
                print(f"Raw response: {len(issues)} issues")
                
                # Convert to simple objects for consistency
                for issue in issues:
                    key = issue.get('key')
                    if key and key not in all_issues_dict:
                        summary = issue.get('fields', {}).get('summary', 'No title')
                        issue_obj = type('Issue', (), {
                            'key': key,
                            'fields': type('Fields', (), {
                                'summary': summary
                            })()
                        })()
                        all_issues_dict[key] = issue_obj
                        # Avoid encoding errors in console output
                        try:
                            print(f"  Added issue: {key} - {summary}")
                        except UnicodeEncodeError:
                            print(f"  Added issue: {key} - [title with non-ASCII chars]")
                
                print(f"Query {i} result: {len(issues)} issues found")
                
            except UnicodeEncodeError as ue:
                logger.warning(f"Encoding error with query {i}, but continuing...")
                print(f"Encoding error with query {i}, but continuing...")
                continue
            except Exception as e:
                logger.error(f"ERROR with JQL query {i}: {e}")
                logger.error(f"Query was: {jql}")
                try:
                    print(f"ERROR with JQL query {i}: {e}")
                    print(f"Query was: {jql}")
                except UnicodeEncodeError:
                    print(f"ERROR with JQL query {i}: [error message with non-ASCII chars]")
                continue
        
        # Convert to list and sort by key
        issues_list = list(all_issues_dict.values())
        issues_list.sort(key=lambda x: x.key)
        
        logger.info(f"Total unique issues found for {calendar.month_name[month_num]} {year}: {len(issues_list)}")
        print(f"Total unique issues found for {calendar.month_name[month_num]} {year}: {len(issues_list)}")
        return issues_list
    
    def generate_report(self, month_text):
        """Generate report for the specified month"""
        try:
            # Parse month and year
            current_year = datetime.now().year
            month_num, year = self.parse_month_year(month_text, current_year)
            month_name = calendar.month_name[month_num]
            
            logger.info(f"Generating report for {month_name} {year}")
            print(f"Generating report for {month_name} {year}")
            
            # Get user tasks
            issues = self.get_user_tasks_for_month(month_num, year)
            
            if not issues:
                logger.warning(f"No tasks found for {month_name} {year}")
                print(f"No tasks found for {month_name} {year}")
                return None
            
            # Load template
            template_path = self.project_root / "resources" / "report_work_template.docx"
            if not template_path.exists():
                raise FileNotFoundError(f"Template not found: {template_path}")
            
            doc = Document(template_path)
            
            # Prepare replacements for template
            replacements = {
                '<MM>': f"{month_num:02d}",
                '<yyyy>': str(year)
            }
            
            # Replace MM and yyyy in document
            self._replace_placeholders(doc, replacements)
            
            # Add tasks to document
            self._add_tasks_to_document(doc, issues)
            
            # Generate output filename (transliterate cyrillic to avoid encoding issues)
            reports_dir = self.project_root / "reports"
            reports_dir.mkdir(exist_ok=True)
            
            # Transliterate cyrillic author name to latin characters for filename
            author_safe = self._transliterate_filename(self.report_author)
            filename = f"{author_safe}_work_report_{month_name.lower()}_{year}.docx"
            output_path = reports_dir / filename
            
            # Save document
            doc.save(output_path)
            logger.info(f"Report generated successfully: {output_path}")
            print(f"Report generated: {output_path}")
            
            return output_path
            
        except Exception as e:
            logger.error(f"Error generating report: {e}")
            print(f"Error generating report: {e}")
            raise
    
    def _replace_placeholders(self, doc, replacements):
        """Replace placeholders in Word document"""
        # Replace in paragraphs
        for paragraph in doc.paragraphs:
            for placeholder, value in replacements.items():
                if placeholder in paragraph.text:
                    paragraph.text = paragraph.text.replace(placeholder, value)
        
        # Replace in tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        for placeholder, value in replacements.items():
                            if placeholder in paragraph.text:
                                paragraph.text = paragraph.text.replace(placeholder, value)
    
    def _add_tasks_to_document(self, doc, issues):
        """Add tasks list to the document by replacing placeholder"""
        # Build tasks text
        tasks_lines = []
        for issue in issues:
            task_key = issue.key
            task_title = issue.fields.summary
            # Format: (<task>)<title_task>
            tasks_lines.append(f"({task_key}){task_title}")
        
        # Join all tasks with newlines
        tasks_text = '\n'.join(tasks_lines)
        
        # Find and replace the placeholder (<task>)<title_task>
        placeholder = '(<task>)<title_task>'
        
        # Search in paragraphs
        for paragraph in doc.paragraphs:
            if placeholder in paragraph.text:
                paragraph.text = paragraph.text.replace(placeholder, tasks_text)
                return
        
        # Search in tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        if placeholder in paragraph.text:
                            paragraph.text = paragraph.text.replace(placeholder, tasks_text)
                            return
        
        # If placeholder not found, add tasks at the end (fallback)
        print(f"Warning: Placeholder '{placeholder}' not found in template, adding tasks at the end")
        doc.add_heading('Tasks:', level=2)
        doc.add_paragraph(tasks_text)
    
    def _transliterate_filename(self, text):
        """Transliterate cyrillic characters to latin for safe filenames"""
        # Simple transliteration map for cyrillic to latin
        translit_map = {
            'а': 'a', 'б': 'b', 'в': 'v', 'г': 'g', 'д': 'd', 'е': 'e', 'ё': 'yo',
            'ж': 'zh', 'з': 'z', 'и': 'i', 'й': 'j', 'к': 'k', 'л': 'l', 'м': 'm',
            'н': 'n', 'о': 'o', 'п': 'p', 'р': 'r', 'с': 's', 'т': 't', 'у': 'u',
            'ф': 'f', 'х': 'h', 'ц': 'c', 'ч': 'ch', 'ш': 'sh', 'щ': 'sch', 'ъ': '',
            'ы': 'y', 'ь': '', 'э': 'e', 'ю': 'yu', 'я': 'ya',
            
            'А': 'A', 'Б': 'B', 'В': 'V', 'Г': 'G', 'Д': 'D', 'Е': 'E', 'Ё': 'Yo',
            'Ж': 'Zh', 'З': 'Z', 'И': 'I', 'Й': 'J', 'К': 'K', 'Л': 'L', 'М': 'M',
            'Н': 'N', 'О': 'O', 'П': 'P', 'Р': 'R', 'С': 'S', 'Т': 'T', 'У': 'U',
            'Ф': 'F', 'Х': 'H', 'Ц': 'C', 'Ч': 'Ch', 'Ш': 'Sh', 'Щ': 'Sch', 'Ъ': '',
            'Ы': 'Y', 'Ь': '', 'Э': 'E', 'Ю': 'Yu', 'Я': 'Ya'
        }
        
        # Transliterate text
        result = ''
        for char in text:
            if char in translit_map:
                result += translit_map[char]
            elif char.isalnum() or char in '-_':
                result += char
            else:
                result += '_'
        
        # Clean up multiple underscores and spaces
        result = result.replace(' ', '_')
        while '__' in result:
            result = result.replace('__', '_')
        
        return result.strip('_')


def main():
    """Main function for command line usage"""
    if len(sys.argv) != 2:
        print("Usage: python jira_report_generator.py <month>")
        print("Example: python jira_report_generator.py september")
        sys.exit(1)
    
    month_text = sys.argv[1]
    
    try:
        generator = JiraReportGenerator()
        report_path = generator.generate_report(month_text)
        
        if report_path:
            print(f"\nSUCCESS Report successfully generated: {report_path}")
        else:
            print("\nINFO No report generated (no tasks found)")
            
    except Exception as e:
        print(f"\nERROR: {e}")
        sys.exit(1)


if __name__ == "__main__":
    main()
