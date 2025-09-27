#!/usr/bin/env python3
"""
Jira Report Generator Script

Generates work reports from Jira tasks for a specific month
"""

import os
import sys
import calendar
from datetime import datetime, timedelta
from pathlib import Path
from atlassian import Jira
from docx import Document
from dotenv import load_dotenv

# Load environment variables
load_dotenv('bot/config.env')

class JiraReportGenerator:
    def __init__(self):
        self.jira_server = os.getenv('JIRA_SERVER')
        self.jira_username = os.getenv('JIRA_USERNAME')
        self.jira_api_token = os.getenv('JIRA_API_TOKEN')
        self.report_author = os.getenv('REPORT_AUTHOR', 'Unknown User')
        self.jira_client = None
        
        if not all([self.jira_server, self.jira_username, self.jira_api_token]):
            raise ValueError("Missing Jira configuration. Check config.env file.")
        
        self._connect_to_jira()
    
    def _connect_to_jira(self):
        """Connect to Jira using API token authentication"""
        try:
            self.jira_client = Jira(
                url=self.jira_server,
                username=self.jira_username,
                password=self.jira_api_token
            )
            print(f"Successfully connected to Jira: {self.jira_server}")
        except Exception as e:
            raise ConnectionError(f"Failed to connect to Jira: {e}")
    
    def parse_month_year(self, month_text, current_year=None):
        """Parse month name and return month number and year"""
        if current_year is None:
            current_year = datetime.now().year
            
        month_names = {
            'january': 1, 'jan': 1, 'janvar': 1,
            'february': 2, 'feb': 2, 'fevral': 2,
            'march': 3, 'mar': 3, 'mart': 3,
            'april': 4, 'apr': 4, 'aprel': 4,
            'may': 5, 'maj': 5,
            'june': 6, 'jun': 6, 'ijun': 6,
            'july': 7, 'jul': 7, 'ijul': 7,
            'august': 8, 'aug': 8, 'avgust': 8,
            'september': 9, 'sep': 9, 'sentjabr': 9,
            'october': 10, 'oct': 10, 'oktjabr': 10,
            'november': 11, 'nov': 11, 'nojabr': 11,
            'december': 12, 'dec': 12, 'dekabr': 12
        }
        
        month_lower = month_text.lower().strip()
        month_num = month_names.get(month_lower)
        
        if month_num is None:
            raise ValueError(f"Unknown month: {month_text}")
        
        return month_num, current_year
    
    def get_month_date_range(self, month_num, year):
        """Get start and end dates for the given month"""
        first_day = datetime(year, month_num, 1)
        last_day_num = calendar.monthrange(year, month_num)[1]
        last_day = datetime(year, month_num, last_day_num, 23, 59, 59)
        
        return first_day, last_day
    
    def get_user_tasks_for_month(self, month_num, year):
        """Get all tasks where current user participated in the given month"""
        first_day, last_day = self.get_month_date_range(month_num, year)
        
        # Format dates for JQL
        start_date = first_day.strftime('%Y-%m-%d')
        end_date = last_day.strftime('%Y-%m-%d')
        
        # JQL to find issues where user was assignee, reporter, or made comments
        jql_queries = [
            f'assignee = currentUser() AND updated >= "{start_date}" AND updated <= "{end_date}"',
            f'reporter = currentUser() AND created >= "{start_date}" AND created <= "{end_date}"',
            f'issueFunction in commented("by currentUser() after {start_date} before {end_date}")'
        ]
        
        all_issues_dict = {}  # Use dict to avoid duplicates by key
        
        for jql in jql_queries:
            try:
                # Use atlassian-python-api's jql method
                issues_data = self.jira_client.jql(jql, limit=200)
                issues = issues_data.get('issues', [])
                
                # Convert to simple objects for consistency
                for issue in issues:
                    key = issue.get('key')
                    if key and key not in all_issues_dict:
                        issue_obj = type('Issue', (), {
                            'key': key,
                            'fields': type('Fields', (), {
                                'summary': issue.get('fields', {}).get('summary', 'No title')
                            })()
                        })()
                        all_issues_dict[key] = issue_obj
                
                print(f"Found {len(issues)} issues with JQL: {jql}")
            except Exception as e:
                print(f"Warning: Error with JQL query '{jql}': {e}")
                continue
        
        # Convert to list and sort by key
        issues_list = list(all_issues_dict.values())
        issues_list.sort(key=lambda x: x.key)
        
        print(f"Total unique issues found for {calendar.month_name[month_num]} {year}: {len(issues_list)}")
        return issues_list
    
    def generate_report(self, month_text):
        """Generate report for the specified month"""
        try:
            # Parse month and year
            current_year = datetime.now().year
            month_num, year = self.parse_month_year(month_text, current_year)
            month_name = calendar.month_name[month_num]
            
            print(f"Generating report for {month_name} {year}")
            
            # Get user tasks
            issues = self.get_user_tasks_for_month(month_num, year)
            
            if not issues:
                print(f"No tasks found for {month_name} {year}")
                return None
            
            # Load template
            template_path = Path("resources/report_work_template.docx")
            if not template_path.exists():
                raise FileNotFoundError(f"Template not found: {template_path}")
            
            doc = Document(template_path)
            
            # Prepare replacements for template
            replacements = {
                'MM': f"{month_num:02d}",
                'yyyy': str(year)
            }
            
            # Replace MM and yyyy in document
            self._replace_placeholders(doc, replacements)
            
            # Add tasks to document
            self._add_tasks_to_document(doc, issues)
            
            # Generate output filename
            reports_dir = Path("reports")
            reports_dir.mkdir(exist_ok=True)
            
            filename = f"{self.report_author.replace(' ', '_')}_work_report_{month_name.lower()}_{year}.docx"
            output_path = reports_dir / filename
            
            # Save document
            doc.save(output_path)
            print(f"Report generated: {output_path}")
            
            return output_path
            
        except Exception as e:
            print(f"Error generating report: {e}")
            raise
    
    def _replace_placeholders(self, doc, replacements):
        """Replace placeholders in Word document"""
        # Replace in paragraphs
        for paragraph in doc.paragraphs:
            for placeholder, value in replacements.items():
                if placeholder in paragraph.text:
                    paragraph.text = paragraph.text.replace(placeholder, value)
        
        # Replace in tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        for placeholder, value in replacements.items():
                            if placeholder in paragraph.text:
                                paragraph.text = paragraph.text.replace(placeholder, value)
    
    def _add_tasks_to_document(self, doc, issues):
        """Add tasks list to the document"""
        # Find a paragraph where we want to insert tasks (after MM and yyyy replacements)
        # We'll add tasks at the end of the document
        
        # Add a heading for tasks
        doc.add_heading('Выполненные задачи:', level=2)
        
        # Add each task
        for issue in issues:
            task_key = issue.key
            task_title = issue.fields.summary
            
            # Format: (<task>)<title_task>
            task_text = f"({task_key}){task_title}"
            doc.add_paragraph(task_text)
        
        # Add summary
        doc.add_paragraph("")
        doc.add_paragraph(f"Всего задач: {len(issues)}")


def main():
    """Main function for command line usage"""
    if len(sys.argv) != 2:
        print("Usage: python jira_report_generator.py <month>")
        print("Example: python jira_report_generator.py september")
        sys.exit(1)
    
    month_text = sys.argv[1]
    
    try:
        generator = JiraReportGenerator()
        report_path = generator.generate_report(month_text)
        
        if report_path:
            print(f"\nSUCCESS Report successfully generated: {report_path}")
        else:
            print("\nINFO No report generated (no tasks found)")
            
    except Exception as e:
        print(f"\nERROR: {e}")
        sys.exit(1)


if __name__ == "__main__":
    main()
